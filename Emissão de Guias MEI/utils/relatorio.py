import logging
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

logger = logging.getLogger(__name__)

COR_SUCESSO  = "C6EFCE"  # verde claro
COR_ERRO     = "FFC7CE"  # vermelho claro
COR_CABEC    = "1F497D"  # azul escuro
COR_TEXTO_CB = "FFFFFF"
COR_DEBITO   = "FFD966"  # amarelo-laranja para débitos em atraso


def gerar_relatorio(resultados: list[dict], pasta: Path, sufixo: str | None = None) -> Path:
    pasta.mkdir(exist_ok=True)
    ts    = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome  = f"relatorio_{sufixo}_{ts}.xlsx" if sufixo else f"relatorio_{ts}.xlsx"
    path  = pasta / nome

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados"

    colunas  = ["CNPJ", "Nome", "Competência", "Status", "Tipo Guia", "Arquivo",
                "Débitos em Atraso", "Observação", "Data/Hora"]
    larguras = [18, 35, 14, 12, 16, 55, 50, 40, 20]

    # Cabeçalho
    for col, (titulo, larg) in enumerate(zip(colunas, larguras), start=1):
        cel = ws.cell(row=1, column=col, value=titulo)
        cel.font      = Font(bold=True, color=COR_TEXTO_CB)
        cel.fill      = PatternFill("solid", fgColor=COR_CABEC)
        cel.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cel.column_letter].width = larg

    # Dados
    col_debito = colunas.index("Débitos em Atraso") + 1  # número da coluna (1-based)

    for i, r in enumerate(resultados, start=2):
        valores = [
            r.get("CNPJ", ""),
            r.get("Nome", ""),
            r.get("Competência", ""),
            r.get("Status", ""),
            r.get("Tipo Guia", ""),
            r.get("Arquivo", ""),
            r.get("Débitos em Atraso", ""),
            r.get("Observação", ""),
            r.get("Data/Hora", ""),
        ]
        for col, valor in enumerate(valores, start=1):
            cel = ws.cell(row=i, column=col, value=valor)

        status     = r.get("Status", "")
        tem_debito = bool(r.get("Débitos em Atraso", ""))
        cor        = COR_SUCESSO if status == "SUCESSO" else COR_ERRO

        for col in range(1, len(colunas) + 1):
            fill_cor = COR_DEBITO if (tem_debito and col == col_debito) else cor
            ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=fill_cor)

    # Resumo
    total    = len(resultados)
    sucessos = sum(1 for r in resultados if r.get("Status") == "SUCESSO")
    erros    = total - sucessos

    devedores = [r for r in resultados if r.get("Débitos em Atraso")]

    ws_res = wb.create_sheet("Resumo")
    ws_res["A1"] = "Total"
    ws_res["B1"] = total
    ws_res["A2"] = "Sucesso"
    ws_res["B2"] = sucessos
    ws_res["A3"] = "Falhas"
    ws_res["B3"] = erros
    ws_res["A4"] = "Com Débitos em Atraso"
    ws_res["B4"] = len(devedores)
    if devedores:
        ws_res["B4"].fill = PatternFill("solid", fgColor=COR_DEBITO)

    # Aba de débitos pendentes
    if devedores:
        ws_deb = wb.create_sheet("Débitos Pendentes")

        cab_deb  = ["CNPJ", "Nome", "Status Guia", "Períodos em Atraso"]
        larg_deb = [18, 35, 14, 70]
        for col, (titulo, larg) in enumerate(zip(cab_deb, larg_deb), start=1):
            cel = ws_deb.cell(row=1, column=col, value=titulo)
            cel.font      = Font(bold=True, color=COR_TEXTO_CB)
            cel.fill      = PatternFill("solid", fgColor="8B0000")  # vermelho escuro
            cel.alignment = Alignment(horizontal="center")
            ws_deb.column_dimensions[cel.column_letter].width = larg

        for i, r in enumerate(devedores, start=2):
            vals = [
                r.get("CNPJ", ""),
                r.get("Nome", ""),
                r.get("Status", ""),
                r.get("Débitos em Atraso", ""),
            ]
            for col, v in enumerate(vals, start=1):
                cel = ws_deb.cell(row=i, column=col, value=v)
                cel.fill = PatternFill("solid", fgColor=COR_DEBITO)

    wb.save(path)
    logger.info(f"Relatório salvo: {path}")
    return path
